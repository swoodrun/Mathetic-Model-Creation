"""Build final design config/mapping/provenance; no fit or optimization."""
from pathlib import Path
import csv
import hashlib
import json
from openpyxl import load_workbook

OUT=Path(__file__).resolve().parent
ROOT=OUT.parents[1]
TEAM=ROOT/'C题组员支援/阶段3'
DATA=ROOT/'C题组员支援/阶段2B/C题_阶段2B_标准数据与复现材料'
cfg=json.loads((TEAM/'stage3_config_v2_review_张.json').read_text(encoding='utf-8'))
cfg['status']='final_design_frozen_not_fitted_or_solved'
cfg['version']='stage3-final-v3'
cfg['time'].update({'interpretation':'right_endpoint_working_assumption_not_officially_resolved','control_timing':'ten_minute_idealized_response_and_end_of_interval_settlement_replay','mapping_file':'C题工作区/阶段3_最终定稿/time_mapping_final.csv'})
cfg['scenarios']['group']='same_historical_day_decision_hour_source_release_hour_and_target_slot'
cfg['scenarios']['invalid_price_rule']='any_nonfinite_or_nonpositive_price_rejects_entire_pool; use_single_point_path_fallback; do_not_selectively_drop_extreme_paths'
cfg['scenarios']['projection_diagnostics']['on_exceed']='reject_entire_pool_and_use_uncorrected_point_path_b0; compare_keep_and_change_normally; retain_current_policy'
cfg['scenarios']['zero_valid_paths']='uncorrected_single_point_path_b0_lambda0_keep_current_policy_and_normal_trigger'
cfg['scenarios']['annual_fallback_decision_fraction_warning']=0.1
cfg['scenarios']['on_high_fallback_fraction']='report_coverage_inadequacy_and_diagnose_before_enhancement_claims; no_retrospective_threshold_tuning'
c=cfg['candidates']
c.pop('quantile_correction_alphas',None)
c['margin_policy']='signed_raw_quantile_correction_no_clipping'
c['free_a_reference']={'virtual_shortfall_fixed_zero':True,'balance':'a+d=N_tilde+c+w','optional_debug_v_must_be_near_zero':True}
c['fixed_a_reference']['free_a_reference_v_near_zero_check']=True
c['reference_release_kwh']=c.pop('b_kwh')
r=cfg['risk']
for k in ['trigger_epsilon_yuan','trigger_sensitivity_yuan_values']:r.pop(k,None)
r['zero_paths']='uncorrected_single_point_path_b0_lambda0_with_current_policy_and_normal_trigger'
r['terminal_band_sensitivity']={'enabled_in_main':False,'required_single_factor_comparison':True,'lower_kwh':5000,'upper_kwh':7000,'formula':'kappa*((5000-E)_+ + (E-7000)_+)','apply_to':['reference_generation','scenario_score']}
r['path_stability_check']={'required_when_positive_lambda_reported':True,'dates':['2025-03-20','2025-06-21','2025-09-23','2025-12-21'],'release_hours':[0,6,12,18],'mode':'fixed_candidate_set_leave_one_complete_day_out_reweight_rescore','separate_optional_mode':'regenerate_candidates_and_rescore'}
cfg['information']['historical_day_path_group']=cfg['scenarios']['group']
cfg['information']['old_forecast_reanchor_at_current_observation']=False
cfg['information']['residual_source_must_match_allowed_forecast_age']=True
cfg['experiments']['matched_state_and_information']['P1']['residual_source']='zero_hour_issue_sliced_to_current_remaining_slots'
cfg['experiments']['E02_decomposition']['point_vs_scenario_same_target_set_and_controller']=True
cfg['innovations']['soft_economic_reserve'].update({'priority':'conditional_after_baseline_and_diagnostics','reserve_reference_b_kwh':600,'value_information':'entire_hourly_mu_vector_computed_at_tau_no_new_tail_truth','suffix_excludes_current_step':True,'candidate_count_separate_from_main':True})
cfg['innovations']['pv_information_shapley'].update({'priority':'first_deepening_after_required_comparisons','controller':'main_hard_reference_b_set_0_600_9600','empty_subset_is_not_P0':True,'residual_source_must_match_allowed_forecast_age':True,'inventory_nu_base':'eta_discharge*median(attachment1_tariff)','compute_budget_fallback':'four_nested_subsets_with_order_dependence_disclosed_no_shapley_claim'})
cfg['innovations']['emergency_diagnostics']={'required':True,'extra_discharge_kwh':'min(h,M-d,eta_discharge*max(E_end-Emin,0))','avoidance_cost':'5*p*extra_discharge','interpretation':'one_step_physical_substitutability_not_free_annual_savings'}
cfg['innovations']['capacity_sensitivity']={'enabled':False,'scope':'Q1_first','capacity_factors':[0.8,1,1.2],'minimum_fraction':0.1,'maximum_fraction':0.9,'initial_and_terminal_fraction':0.5,'power_limit_kw_fixed':5000,'capital_payback_not_estimated':True}
cfg['innovations']['price_volatility_sensitivity']={'enabled':False,'gamma':[0.5,1,1.5],'formula':'daily_arithmetic_mean_p * exp(gamma*(log_p-mean_log_p))/daily_mean_exp','full_day_statistics_used_only_for_offline_world_generation':True,'rebuild_causal_forecasts_and_archives_per_world':True}
cfg['implementation']['environment_status']='not_prepared_by_final_design_task'
cfg['implementation']['actual_optimization_run']=False
cfg['implementation']['actual_backtest_run']=False
cfg['export']={'source_workbooks_unchanged':True,'generated_workbook_labels':'canonical_natural_day_from_time_mapping_final','preserve_sheet_structure_and_slot_order':True,'adjustment_table_quantity':'final_total_a','net_adjustment_ledger':'a-g','planned_daily_fee':'sum(p*g)','adjusted_daily_fee':'sum(F(g,a,p))','actual_total_fee_ledger':'sum(F+5*p*h)','evaluation_days':334,'slots_per_day':144,'storage_blocks_per_day':6,'ellipsis_rows':'expand_to_all_days','emergency_events':'merge_consecutive_positive_intervals_within_same_day'}
cfg['data']={'actual':'C题组员支援/阶段2B/C题_阶段2B_标准数据与复现材料/data/actual_10min.csv','q1':'C题组员支援/阶段2B/C题_阶段2B_标准数据与复现材料/data/q1_inputs.csv','pv_forecast':'C题组员支援/阶段2B/C题_阶段2B_标准数据与复现材料/data/pv_forecast_10min.csv','dictionary':'C题工作区/阶段2AB_独立复核/数据字典_复核修订版.md'}
cfg['provenance']={'assembled_from':'stage3_config_v2_review_张.json plus current final decisions','canonical_document':'C题最终方案_阶段3定稿.md','source_manifest':'source_manifest.json','design_contract_only':True,'not_executable_algorithm':True}
(OUT/'stage3_final_config.json').write_text(json.dumps(cfg,ensure_ascii=False,indent=2),encoding='utf-8')

with (DATA/'data/time_mapping.csv').open(encoding='utf-8-sig',newline='') as f: mapping=list(csv.DictReader(f))
wb=load_workbook(ROOT/'2026年数模题目和论文格式规范/C题/附件/附件5/result1.xlsx',read_only=True,data_only=True)
q1labels=[r[0] for r in list(wb['计划购电量'].values)[1:145]];wb.close()
from openpyxl.utils import get_column_letter
for i,row in enumerate(mapping):
    row['original_q1_label']=q1labels[i]
    row['q1_template_row']=i+2
    row['annual_template_column']=get_column_letter(i+2)
    row['export_label']=row['canonical_label']
    row['interpretation']='right_endpoint_assumption_not_official_correction'
with (OUT/'time_mapping_final.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(mapping[0]));w.writeheader();w.writerows(mapping)
sources=list(TEAM.glob('*'))+[ROOT/'C题工作区/阶段3_模型方案设计/阶段3_模型方案设计.md',ROOT/'C题工作区/阶段3_模型方案设计/stage3_config.json',ROOT/'C题工作区/阶段2AB_独立复核/independent_review.json',DATA/'data/time_mapping.csv',ROOT/'2026年数模题目和论文格式规范/C题/C题.pdf',ROOT/'C题最终方案_阶段3定稿.md']
manifest=[{'file':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sources if p.is_file()]
(OUT/'source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print('Final config, 144-slot export mapping and source manifest generated.')
